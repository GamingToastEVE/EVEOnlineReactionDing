"""Checks the reaction formulas of the "Operation Moonshine" workbook against the website.

The workbook hard-codes reaction recipes in its formulas. This module reads every
such formula (needs ``openpyxl``), reduces it to "input quantity per run" and
"output per run", and compares that with the base recipe the calculator API
returns for the same reaction (requested without rigs, so no material bonus).

Recognised sheet layouts (found by their headers, not by sheet name):

* profit sheet   - "Reaction Price" columns, e.g. ``=(100*A4+100*A5+5*A32)/200``
* quantity sheet - "production order" / "Reaction Formulas" planning tables
* formula blocks - "<Reaction> Formula" blocks with ingredient lists
"""

import re
from collections import defaultdict
from dataclasses import dataclass

from openpyxl.utils import column_index_from_string, get_column_letter

from . import settings as settings_mod
from .engine import calculate

FUEL_BLOCKS = ("Helium Fuel Block", "Hydrogen Fuel Block", "Nitrogen Fuel Block", "Oxygen Fuel Block")
RECIPE_GROUPS = ("simple", "complex", "hybrid")
RUNS_PER_JOB = 544


@dataclass
class Finding:
    sheet: str
    cell: str
    reaction: str
    message: str
    fix: str = ""


def load_recipes(client=None):
    """Base recipe per reaction name: {"out": per run, "in": {name: per run}, "group": ...}."""
    values = settings_mod.normalize({"rigs": 0, "facility": "medium"})
    recipes, errors = {}, []
    for row in calculate(values, list(RECIPE_GROUPS), client=client):
        if not row.ok:
            errors.append(f"{row.name}: {row.error}")
            continue
        res, runs = row.result, row.result["runs"]
        recipes[row.name] = {
            "group": row.group.key,
            "out": res["output"]["quantity"] / runs,
            "in": {x["name"]: x["quantity"] / runs for x in res["input"]},
        }
    return recipes, errors


def _fmt(value):
    if value is None:
        return "missing"
    return f"{value:g}" if isinstance(value, (int, float)) else str(value)


def _compare(found, expected):
    """Differences between two {material: qty} dicts as readable strings."""
    out = []
    for name in sorted(set(found) | set(expected)):
        a, b = found.get(name), expected.get(name)
        if a is None or b is None or abs(a - b) > 1e-9:
            out.append(f"{name}: sheet {_fmt(a)}, website {_fmt(b) if b is not None else 'not in recipe'}")
    return out


class Sheet:
    def __init__(self, ws):
        self.ws = ws
        self.title = ws.title
        self.cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    self.cells[c.coordinate] = c.value

    def v(self, ref):
        return self.cells.get(ref.replace("$", ""))

    def text(self, ref):
        value = self.v(ref)
        return value.strip() if isinstance(value, str) else value

    def find_row(self, row, text):
        """Column letters in ``row`` whose value equals ``text``."""
        hits = []
        for ref, value in self.cells.items():
            m = re.fullmatch(r"([A-Z]+)(\d+)", ref)
            if int(m.group(2)) == row and isinstance(value, str) and value.strip() == text:
                hits.append(m.group(1))
        return sorted(hits, key=column_index_from_string)


# --- profit sheet ("Reaction Price") -----------------------------------------

_TERM = re.compile(r"(?:(\d+(?:\.\d+)?)\*)?\$?([A-Z]+)\$?(\d+)(?:\*(\d+(?:\.\d+)?))?")


def _parse_recipe_formula(formula):
    m = re.fullmatch(r"=\((.*)\)/(\d+(?:\.\d+)?)", formula.replace(" ", ""))
    if not m:
        return None
    terms = []
    for part in m.group(1).split("+"):
        t = _TERM.fullmatch(part)
        if not t:
            return None
        coef = float(t.group(1) or t.group(4) or 1)
        terms.append((coef, t.group(2), int(t.group(3))))
    return terms, float(m.group(2))


def check_profit_sheet(sheet, recipes):
    """Sheet with "(qty*price + ...)/output" formulas (e.g. '7 - Reactions Calculations Prof')."""
    findings = []
    # every recipe formula sits left of the column holding the reaction name
    for ref, formula in sheet.cells.items():
        if not isinstance(formula, str) or not formula.startswith("=("):
            continue
        parsed = _parse_recipe_formula(formula)
        if not parsed:
            continue
        col, row = re.fullmatch(r"([A-Z]+)(\d+)", ref).groups()
        idx = column_index_from_string(col)
        product = next((sheet.text(f"{get_column_letter(idx + k)}{row}") for k in (2, 1, 3)
                        if sheet.text(f"{get_column_letter(idx + k)}{row}") in recipes), None)
        if product is None:
            continue
        recipe = recipes[product]
        terms, divisor = parsed
        found, price_cols = {}, set()
        for coef, tcol, trow in terms:
            name = _name_for_price_ref(sheet, tcol, trow, recipes)
            found[name] = found.get(name, 0) + coef
            price_cols.add((tcol, name))
        msgs = _compare(found, recipe["in"])
        if divisor != recipe["out"]:
            msgs.append(f"divisor (output per run): sheet {divisor:g}, website {recipe['out']:g}")
        for msg in msgs:
            findings.append(Finding(sheet.title, ref, product, msg,
                                    _profit_fix(sheet, recipe, terms, divisor, recipes)))
        cols = {c for c, n in price_cols if n not in recipes}  # raw material price columns
        if len(cols) > 1:
            findings.append(Finding(sheet.title, ref, product,
                                    f"mixes price columns {', '.join(sorted(cols))} for raw materials"))
    return findings


def _name_for_price_ref(sheet, col, row, recipes):
    """Name of the item a price cell belongs to: the first text cell to its right."""
    idx = column_index_from_string(col)
    for k in range(1, 6):
        value = sheet.text(f"{get_column_letter(idx + k)}{row}")
        if isinstance(value, str) and not value.startswith("="):
            return value
    return f"?{col}{row}"


def _profit_fix(sheet, recipe, terms, divisor, recipes):
    """Rebuild the formula with the website quantities, keeping the sheet's own cell refs."""
    refs = {}
    for _, tcol, trow in terms:
        refs.setdefault(_name_for_price_ref(sheet, tcol, trow, recipes), f"{tcol}{trow}")
    parts = []
    for name, qty in recipe["in"].items():
        ref = refs.get(name)
        if ref is None:
            return ""
        parts.append(f"{qty:g}*{ref}")
    return f"=({'+'.join(parts)})/{recipe['out']:g}"


# --- quantity planning sheets ------------------------------------------------

def _layout(sheet):
    """Column letters of a planning sheet, located via the header row (row 2)."""
    def first(text, row=2):
        hits = sheet.find_row(row, text)
        return hits[0] if hits else None

    order = first("production order")
    reaction_formulas = first("Reaction Formulas")
    if not order or not reaction_formulas:
        return None
    demand = first("Demand")
    goo = next(c for c in sheet.find_row(3, "Atmospheric Gases")
               if column_index_from_string(c) > column_index_from_string(demand))
    simple = sheet.find_row(3, "Caesarium Cadmide")[0]
    complex_ = next(c for c in sheet.find_row(3, "Crystalline Carbonide")
                    if column_index_from_string(c) > column_index_from_string(reaction_formulas))
    rig_col = first("Rig", row=1)
    return {
        "goo": goo, "demand": demand, "simple": simple, "order": order,
        "batch": first("batch production"),
        "stock": get_column_letter(column_index_from_string(order) - 1),
        "total": reaction_formulas, "complex": complex_, "out": first("1 run output"),
        "rig": f"{get_column_letter(column_index_from_string(rig_col) + 1)}1" if rig_col else None,
    }


def check_quantity_sheet(sheet, recipes):
    L = _layout(sheet)
    if not L:
        return []
    findings = []
    simples = {r: sheet.text(f"{L['simple']}{r}") for r in range(3, 31)
               if recipes.get(sheet.text(f"{L['simple']}{r}"), {}).get("group") == "simple"}
    complexes = {r: sheet.text(f"{L['complex']}{r}") for r in range(3, 25)
                 if recipes.get(sheet.text(f"{L['complex']}{r}"), {}).get("group") == "complex"}
    simple_row = {n: r for r, n in simples.items()}

    # complex output per run
    for r, name in complexes.items():
        out = sheet.v(f"{L['out']}{r}")
        if out != recipes[name]["out"]:
            findings.append(Finding(sheet.title, f"{L['out']}{r}", name,
                                    f"output per run: sheet {_fmt(out)}, website {recipes[name]['out']:g}",
                                    f"{recipes[name]['out']:g}"))

    # simple reactions consumed by each complex job (production order column)
    used = defaultdict(dict)
    for r, sname in simples.items():
        formula = str(sheet.v(f"{L['order']}{r}") or "")
        for m in re.finditer(r"ROUNDUP\((\d+)\*\(1\+\$?[A-Z]+\$?1\),0\)\*\$?" + L["total"] + r"(\d+)",
                             formula, re.I):
            used[complexes.get(int(m.group(2)), f"row {m.group(2)}")][sname] = int(m.group(1)) / RUNS_PER_JOB
        for m in re.finditer(r"\$?" + L["total"] + r"(\d+)\*" + str(RUNS_PER_JOB), formula):
            used[complexes.get(int(m.group(1)), f"row {m.group(1)}")][sname] = 1.0
        batch = re.search(r"/(\d+),0\)\*\d+", str(sheet.v(f"{L['batch']}{r}") or ""))
        expected = RUNS_PER_JOB * recipes[sname]["out"]
        if batch and int(batch.group(1)) != expected:
            findings.append(Finding(sheet.title, f"{L['batch']}{r}", sname,
                                    f"batch size {batch.group(1)} != {RUNS_PER_JOB} runs x "
                                    f"{recipes[sname]['out']:g}"))
    for r, cname in complexes.items():
        expected = {k: q for k, q in recipes[cname]["in"].items() if k not in FUEL_BLOCKS}
        for msg in _compare(used.get(cname, {}), expected):
            findings.append(Finding(sheet.title, f"{L['order']} column", cname, "per run " + msg))

    # raw material (moon goo) and fuel demand
    rig = sheet.v(L["rig"]) if L["rig"] else 0
    for r in range(3, 40):
        material = sheet.text(f"{L['goo']}{r}")
        formula = sheet.v(f"{L['demand']}{r}")
        if not isinstance(material, str) or not isinstance(formula, str) or not formula.startswith("="):
            continue
        cell = f"{L['demand']}{r}"
        if material in FUEL_BLOCKS:
            findings += _check_fuel(sheet, L, cell, material, formula, simples, complexes, recipes)
            continue
        found = {}
        for m in re.finditer(r"MAX\(\(?\$?" + L["order"] + r"(\d+)-\$?" + L["stock"] + r"(\d+)\)?"
                             r"(\*\((\d+)/(\d+)\))?", formula, re.I):
            coef = int(m.group(4)) / int(m.group(5)) if m.group(3) else 1.0
            found[simples.get(int(m.group(1)), f"row {m.group(1)}")] = coef
        expected = {n: recipes[n]["in"][material] / recipes[n]["out"]
                    for n in simples.values() if material in recipes[n]["in"]}
        diffs = _compare(found, expected)
        if diffs:
            jobs = L["batch"] and get_column_letter(column_index_from_string(L["batch"]) + 1)
            parts = []
            for n, per_unit in expected.items():
                qty = per_unit * recipes[n]["out"] * RUNS_PER_JOB
                parts.append(f"ROUNDUP({qty:.0f}*(1+${L['rig'][:-1]}$1),0)*{jobs}{simple_row[n]}")
            fix = "=" + "+".join(parts)
            for i, d in enumerate(diffs):
                findings.append(Finding(sheet.title, cell, material, f"per unit of product: {d}",
                                        fix if i == len(diffs) - 1 else ""))
        if re.search(r"\d+/\d+\)\*\(1\+", formula) and rig:
            m = re.search(r"\((\d+)/(\d+)\)", formula)
            if m and abs(int(m.group(1)) / int(m.group(2)) - 200 * (1 + rig)) < 0.5:
                findings.append(Finding(sheet.title, cell, material,
                                        f"factor {m.group(1)}/{m.group(2)} already contains the rig bonus "
                                        "and is multiplied by (1+rig) again"))

    findings += _check_hybrid_block(sheet, recipes)
    return findings


def _check_fuel(sheet, L, cell, fuel, formula, simples, complexes, recipes):
    findings = []
    found = {}
    for m in re.finditer(r"ROUNDUP\((\d+)/(\d+)\*\$?" + L["batch"] + r"(\d+)", formula, re.I):
        found[simples.get(int(m.group(3)), f"row {m.group(3)}")] = int(m.group(1)) / int(m.group(2))
    expected = {n: recipes[n]["in"][fuel] / recipes[n]["out"]
                for n in simples.values() if fuel in recipes[n]["in"]}
    for d in _compare(found, expected):
        findings.append(Finding(sheet.title, cell, fuel, f"simple reactions, per unit of product: {d}"))
    found_c = {complexes.get(int(r), f"row {r}") for r in
               re.findall(r"\d\*\$?" + L["total"] + r"(\d+)\*", formula)}
    expected_c = {n for n in complexes.values() if fuel in recipes[n]["in"]}
    for n in sorted(found_c - expected_c):
        findings.append(Finding(sheet.title, cell, fuel, f"complex reaction {n} does not use {fuel}"))
    for n in sorted(expected_c - found_c):
        findings.append(Finding(sheet.title, cell, fuel, f"complex reaction {n} needs {fuel} but is missing"))
    return findings


def _check_hybrid_block(sheet, recipes):
    """Hybrid planning block: input demand = ROUNDUP((N<row>*qty)*(1+rig)) per hybrid job row."""
    findings = []
    products = {}
    for ref, value in sheet.cells.items():
        if isinstance(value, str) and recipes.get(value.strip(), {}).get("group") == "hybrid":
            col, row = re.fullmatch(r"([A-Z]+)(\d+)", ref).groups()
            products.setdefault(col, {})[int(row)] = value.strip()
    # demand formulas reference the job rows of the hybrid table, e.g. roundup((N41*100)*(1+AC1),0)
    used = defaultdict(lambda: defaultdict(dict))
    fuel_cells = defaultdict(dict)
    for ref, formula in sheet.cells.items():
        if not isinstance(formula, str):
            continue
        hits = re.findall(r"\(\(\$?([A-Z]+)(\d+)\*(\d+)\)", formula)
        if not hits:
            continue
        col, row = re.fullmatch(r"([A-Z]+)(\d+)", ref).groups()
        material = _name_for_price_ref(sheet, col, int(row), recipes)
        for jcol, jrow, qty in hits:
            if material in FUEL_BLOCKS:
                fuel_cells[(ref, material)][int(jrow)] = float(qty)
            else:
                used[jcol][int(jrow)][material] = float(qty)
    for jcol, rows in used.items():
        names = next((n for n in products.values() if set(rows) <= set(n)), None)
        if names is None:
            continue
        for r, name in sorted(names.items()):
            expected = {k: q for k, q in recipes[name]["in"].items() if k not in FUEL_BLOCKS}
            for d in _compare(rows.get(r, {}), expected):
                findings.append(Finding(sheet.title, f"{jcol}{r} refs", name, "per run " + d))
            out_col = get_column_letter(column_index_from_string(jcol) + 2)
            out = sheet.v(f"{out_col}{r}")
            if isinstance(out, (int, float)) and out != recipes[name]["out"]:
                findings.append(Finding(sheet.title, f"{out_col}{r}", name,
                                        f"output per run: sheet {out:g}, website {recipes[name]['out']:g}"))
        for (cell, fuel), job_rows in sorted(fuel_cells.items()):
            found = {names[r] for r in job_rows if r in names}
            expected = {n for n in names.values() if fuel in recipes[n]["in"]}
            for n in sorted(found - expected):
                need = next(f for f in FUEL_BLOCKS if f in recipes[n]["in"])
                findings.append(Finding(sheet.title, cell, n, f"hybrid counted with {fuel}, website uses {need}"))
            for n in sorted(expected - found):
                findings.append(Finding(sheet.title, cell, n, f"{fuel} for this hybrid is missing"))
    return findings


# --- "<Reaction> Formula" blocks ----------------------------------------------

def check_formula_blocks(sheet, recipes):
    findings = []
    for ref, value in sheet.cells.items():
        if not isinstance(value, str) or not value.strip().endswith("Formula"):
            continue
        product = value.strip()[: -len("Formula")].strip()
        if product not in recipes:
            continue
        col, row = re.fullmatch(r"([A-Z]+)(\d+)", ref).groups()
        qty_col = get_column_letter(column_index_from_string(col) + 1)
        found, out, r = {}, None, int(row) + 1
        while r < int(row) + 12:
            name = sheet.text(f"{col}{r}")
            if name == product:
                out = sheet.v(f"{qty_col}{r}")
                break
            if isinstance(name, str) and isinstance(sheet.v(f"{qty_col}{r}"), (int, float)):
                found[name] = float(sheet.v(f"{qty_col}{r}"))
            r += 1
        msgs = _compare(found, recipes[product]["in"])
        if out != recipes[product]["out"]:
            msgs.append(f"output per run: sheet {_fmt(out)}, website {recipes[product]['out']:g}")
        for msg in msgs:
            findings.append(Finding(sheet.title, ref, product, msg))
    return findings


def check_workbook(path, client=None):
    import openpyxl

    recipes, errors = load_recipes(client)
    wb = openpyxl.load_workbook(path)
    report = {"sheets": {}, "errors": errors}
    for ws in wb.worksheets:
        sheet = Sheet(ws)
        findings = check_profit_sheet(sheet, recipes) + check_quantity_sheet(sheet, recipes) \
            + check_formula_blocks(sheet, recipes)
        if findings or any(isinstance(v, str) and v.strip() in recipes for v in sheet.cells.values()):
            report["sheets"][ws.title] = findings
    return report
